"""Synthetic-only contract tests for the provisional Portland queue.

All database access uses clean_state's temporary database. No operator cache,
credentials, customer text, or live retrieval is needed by this module.
"""
import copy
import io
import re
import sqlite3
from html import unescape
from urllib.parse import parse_qs, urlsplit

import pytest
from flask import template_rendered
from openpyxl import load_workbook
from werkzeug.datastructures import MultiDict

import app


GRAMMAR = "ascii-10-comma-v1"
ACTIVATION = "FRESHDESK_PORTLAND_APPROVED_GRAMMAR"
PORTLAND = "7" + "0" * 9
PORTLAND_OTHER = "7" + "1" * 9
OTHER = "8" + "0" * 9
STAMP = "2026-08-05T12:00:00Z"
OLDER = "2026-08-01T12:00:00Z"


@pytest.fixture(autouse=True)
def portland_isolation(clean_state, monkeypatch):
    """Operator activation/offline-cache settings cannot affect these tests."""
    monkeypatch.delenv(ACTIVATION, raising=False)
    monkeypatch.delenv("FRESHDESK_OFFLINE_CACHE", raising=False)
    monkeypatch.setenv("FRESHDESK_OFFLINE_CACHE", "1")
    monkeypatch.setenv("FRESHDESK_OFFLINE", "1")


def ticket(ticket_id=960001, order=PORTLAND, **changes):
    value = {
        "id": ticket_id,
        "subject": "Photo/Video Request synthetic test",
        "status": 2,
        "type": "Guest Callback/Follow-Up",
        "group_id": 154000437139,
        "custom_fields": {"cf_order": order},
        "tags": ["PHOTOS"],
        "priority": 2,
        "created_at": OLDER,
        "updated_at": STAMP,
    }
    value.update(changes)
    return value


def config(scope="main", **changes):
    # Keep workflow routing enabled rather than using Show All's bypass.
    return dict(app.DEFAULT_FILTERS, queue_scope=scope, **changes)


def ids(view):
    return [row["id"] for row in view[0]]


def html_ids(response):
    return [int(value) for value in re.findall(rb'<tr[^>]*data-ticket-id="(\d+)"', response.data)]


def activate(monkeypatch, enabled):
    if enabled:
        monkeypatch.setenv(ACTIVATION, GRAMMAR)
    else:
        monkeypatch.delenv(ACTIVATION, raising=False)


def install_pool(monkeypatch, pool):
    monkeypatch.setattr(app, "get_ticket_pool", lambda: (copy.deepcopy(pool), None))


def snapshot(db_path):
    with sqlite3.connect(db_path) as db:
        return tuple(
            tuple(db.execute(f"SELECT * FROM {table} ORDER BY ticket_id"))
            for table in ("review_state", "closed_review_state")
        )


def seed(db_path, ticket_id, result, reviewed_at=STAMP):
    with sqlite3.connect(db_path) as db:
        db.execute(
            "INSERT INTO review_state (ticket_id, review_result, first_opened_at, "
            "last_opened_at, last_review_change_at, reviewed_updated_at, note, "
            "created_at, modified_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (ticket_id, result, OLDER, STAMP, OLDER, reviewed_at,
             "synthetic preserved note", OLDER, STAMP),
        )
        db.execute(
            "INSERT INTO closed_review_state (ticket_id, review_result, "
            "first_opened_at, last_opened_at, last_review_change_at, "
            "reviewed_updated_at, closed_at_snapshot, note, created_at, modified_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (ticket_id, result, OLDER, STAMP, OLDER, reviewed_at, OLDER,
             "synthetic closed note", OLDER, STAMP),
        )


def test_grammar_identifier():
    assert app.PORTLAND_ORDER_GRAMMAR == GRAMMAR


@pytest.mark.parametrize("raw, expected", [
    (PORTLAND, (PORTLAND,)),
    (" \t" + PORTLAND + "\t ", (PORTLAND,)),
    (PORTLAND + "," + OTHER, (PORTLAND, OTHER)),
    (PORTLAND + " \t,\t " + PORTLAND_OTHER, (PORTLAND, PORTLAND_OTHER)),
    (PORTLAND + "," + PORTLAND + "," + OTHER, (PORTLAND, OTHER)),
    ("0" * 10, ("0" * 10,)),
])
def test_supported_grammar_and_distinct_reference_order(raw, expected):
    assert app.parse_sales_orders(ticket(order=raw)) == expected


@pytest.mark.parametrize("raw", [
    None, True, 7 * 10**9, 7.0, [], [PORTLAND], {}, {"order": PORTLAND},
    "", " \t", "7" * 9, "7" * 11,
    PORTLAND + "\n", "\n" + PORTLAND, PORTLAND + "\r\n",
    PORTLAND + ",\n" + OTHER, PORTLAND + "\v", PORTLAND + "\f",
    "\u00a0" + PORTLAND, PORTLAND + "\u200b", "７" * 10, "٧" * 10,
    "Order " + PORTLAND, "SO:" + PORTLAND, "#" + PORTLAND,
    PORTLAND + ";" + OTHER, PORTLAND + "/" + OTHER,
    PORTLAND + " " + OTHER, PORTLAND + ",," + OTHER,
    "," + PORTLAND, PORTLAND + ",", PORTLAND + ".", "-" + PORTLAND,
])
def test_unsupported_entire_field_is_unknown_not_partially_parsed(raw):
    value = ticket(order=raw)
    assert app.parse_sales_orders(value) == ()
    assert app.classify_portland_ticket(value) == "unknown"


@pytest.mark.parametrize("fields", [None, [], "invalid", {}, {"other": PORTLAND}])
def test_missing_or_malformed_custom_fields(fields):
    assert app.parse_sales_orders(ticket(custom_fields=fields)) == ()
    assert app.classify_portland_ticket(ticket(custom_fields=fields)) == "unknown"


@pytest.mark.parametrize("raw, expected", [
    (PORTLAND, "portland_only"),
    (PORTLAND + "," + PORTLAND_OTHER, "portland_only"),
    (OTHER, "non_portland_only"),
    (OTHER + "," + "9" * 10, "non_portland_only"),
    (PORTLAND + "," + OTHER, "mixed"),
    (OTHER + "," + PORTLAND, "mixed"),
    (None, "unknown"),
])
def test_classification_is_provisional_and_independent_of_activation(monkeypatch, raw, expected):
    for enabled in (False, True):
        activate(monkeypatch, enabled)
        assert app.classify_portland_ticket(ticket(order=raw)) == expected


def test_pure_helpers_do_not_mutate_ticket_or_consult_activation(monkeypatch):
    value = ticket(order=PORTLAND + "," + PORTLAND)
    before = copy.deepcopy(value)

    def forbidden():
        pytest.fail("pure classification consulted routing activation")

    monkeypatch.setattr(app, "portland_routing_enabled", forbidden)
    assert app.parse_sales_orders(value) == (PORTLAND,)
    assert app.classify_portland_ticket(value) == "portland_only"
    assert app.queue_scope_for_ticket(value) == "main"
    assert app.queue_scope_for_ticket(value, portland_enabled=True) == "portland"
    assert value == before


def test_order_like_text_elsewhere_is_ignored():
    value = ticket(order=None, subject=PORTLAND, tags=[PORTLAND],
                   description=PORTLAND, description_text=PORTLAND,
                   custom_fields={"cf_other": PORTLAND},
                   requester={"name": PORTLAND}, conversations=[{"body": PORTLAND}])
    assert app.parse_sales_orders(value) == ()
    assert app.classify_portland_ticket(value) == "unknown"


@pytest.mark.parametrize("raw", [None, "", "1", "true", "approved", GRAMMAR.upper(),
                                      " " + GRAMMAR, GRAMMAR + " ", GRAMMAR + "\n", GRAMMAR])
def test_activation_requires_exact_environment_value(monkeypatch, raw):
    if raw is not None:
        monkeypatch.setenv(ACTIVATION, raw)
    assert app.portland_routing_enabled() is (raw == GRAMMAR)


@pytest.mark.parametrize("raw", [OTHER, PORTLAND + "," + OTHER, None, "Order " + PORTLAND])
@pytest.mark.parametrize("triage", [False, True])
def test_only_portland_only_can_override_existing_scope(raw, triage):
    value = ticket(order=raw, tags=[] if triage else ["PHOTOS"])
    expected = "triage" if triage else "main"
    assert app.queue_scope_for_ticket(value) == expected
    assert app.queue_scope_for_ticket(value, portland_enabled=True) == expected


@pytest.mark.parametrize("triage", [False, True])
def test_portland_override_requires_explicit_enabled_argument(triage):
    value = ticket(tags=[] if triage else ["PHOTOS"])
    assert app.queue_scope_for_ticket(value) == ("triage" if triage else "main")
    assert app.queue_scope_for_ticket(value, portland_enabled=True) == "portland"


@pytest.fixture
def partition_pool():
    return [ticket(960001), ticket(960002, tags=[]),
            ticket(960003, order=OTHER), ticket(960004, order=None, tags=[]),
            ticket(960005, order=PORTLAND + "," + OTHER)]


def test_preview_overlaps_without_excluding_main_or_triage(partition_pool):
    expected = {"main": [960001, 960003, 960005], "triage": [960002, 960004],
                "portland": [960001, 960002]}
    for scope, wanted in expected.items():
        view = app.build_current_queue_view(partition_pool, config(scope))
        assert len(view) == 7
        assert ids(view) == wanted
        assert view[2] == {"main": 3, "triage": 2, "portland": 0}
        assert view[6] == 2
        assert view[3] is True
        assert view[4] is None
        assert view[5] is False


def test_activated_memberships_and_counts_are_disjoint(monkeypatch, partition_pool):
    activate(monkeypatch, True)
    expected = {"main": [960003, 960005], "triage": [960004], "portland": [960001, 960002]}
    memberships = []
    for scope, wanted in expected.items():
        view = app.build_current_queue_view(partition_pool, config(scope))
        assert ids(view) == wanted
        assert view[2] == {"main": 2, "triage": 1, "portland": 2}
        memberships.append(set(ids(view)))
    assert sum(map(len, memberships)) == len(set.union(*memberships)) == len(partition_pool)


@pytest.mark.parametrize("enabled", [False, True])
@pytest.mark.parametrize("result, destination", [
    ("Unreviewed", "main"), ("Opened / In Review", "main"),
    ("Needs Supervisor Review", "supervisor"), ("Needs Follow-Up", "followup"),
    ("Resolved", "resolved"), ("No Action", "no_action"),
    ("Not Applicable to Me", "no_action"), ("No Action Needed", "no_action"),
])
def test_workflow_routing_preserves_every_stored_column(clean_state, monkeypatch, enabled, result, destination):
    activate(monkeypatch, enabled)
    seed(clean_state, 960001, result)
    before = snapshot(clean_state)
    for tab in app.WORKFLOW_TABS:
        view = app.build_current_queue_view([ticket()], config("portland", workflow_tab=tab))
        assert ids(view) == ([960001] if tab == destination else [])
        assert view[1][destination] == 1
        assert sum(view[1].values()) == 1
        if view[0]:
            assert view[0][0]["result"] == ("No Action" if result in app.LEGACY_NO_ACTION_STATES else result)
    assert snapshot(clean_state) == before


@pytest.mark.parametrize("result, tab", [("Resolved", "main"), ("Needs Follow-Up", "main"),
                                         ("No Action", "no_action")])
def test_updated_since_review_behavior_is_reused(clean_state, result, tab):
    seed(clean_state, 960001, result, reviewed_at=OLDER)
    before = snapshot(clean_state)
    row = app.build_current_queue_view([ticket()], config("portland", workflow_tab=tab))[0][0]
    assert row["updated_flag"] is True
    assert row["can_acknowledge"] is True
    assert snapshot(clean_state) == before


@pytest.mark.parametrize("enabled", [False, True])
def test_review_scope_filters_and_deduplication_are_reused(monkeypatch, enabled):
    activate(monkeypatch, enabled)
    reviewed_tag = next(iter(app.REVIEWED_EXCLUSION_TAGS))
    pool = [ticket(), ticket(), ticket(960002, subject="Unrelated synthetic subject", tags=[]),
            ticket(960003, tags=["PHOTOS", reviewed_tag]), ticket(960004, status=5)]
    view = app.build_current_queue_view(pool, config("portland"))
    assert ids(view) == [960001]
    assert sum(view[2].values()) == 1
    assert not app.build_current_queue_view([ticket()], config("portland", missing_tags=True))[0]


def test_url_parser_and_serialization_preserve_portland_with_multidict():
    args = MultiDict([("queue_scope", "main"), ("queue_scope", "portland"),
                      ("workflow_tab", "followup"), ("photo_video_only", "0")])
    parsed = app.filters_from_args(args)
    assert parsed["queue_scope"] == "portland"
    encoded = app.filter_query_string(parsed)
    assert parse_qs(encoded)["queue_scope"] == ["portland"]
    roundtrip = app.filters_from_args(MultiDict((key, value) for key, values in parse_qs(encoded).items() for value in values))
    assert roundtrip == parsed
    assert app.filters_from_args(MultiDict([("queue_scope", "portland"), ("queue_scope", "invalid")])).get("queue_scope", "main") == "main"


def test_url_parameters_cannot_activate_routing(client, monkeypatch, partition_pool):
    install_pool(monkeypatch, partition_pool)
    query = {"queue_scope": "main", ACTIVATION: GRAMMAR,
             "portland_enabled": "1", "portland_routing_enabled": "1"}
    response = client.get("/queue", query_string=query)
    assert response.status_code == 200
    assert html_ids(response) == [960001, 960003, 960005]
    assert app.portland_routing_enabled() is False


@pytest.mark.parametrize("enabled", [False, True])
def test_portland_html_navigation_export_and_forms_preserve_scope(client, monkeypatch, enabled):
    activate(monkeypatch, enabled)
    install_pool(monkeypatch, [ticket()])
    response = client.get("/queue?queue_scope=portland")
    assert response.status_code == 200
    text = response.get_data(as_text=True)
    assert "Portland / No Coverage" in text
    assert html_ids(response) == [960001]
    if not enabled:
        assert "provisional" in text.lower()
        assert "preview" in text.lower()
    links = [unescape(value) for value in re.findall(r'href="([^"]+)"', text)]
    export = next(value for value in links if value.startswith("/queue/export.xlsx?"))
    assert parse_qs(urlsplit(export).query)["queue_scope"] == ["portland"]
    assert re.search(r'name="?queue_scope"?[^>]*value="portland"', text)
    workflow_links = [value for value in links if value.startswith("/queue?") and "workflow_tab=" in value
                      and parse_qs(urlsplit(value).query).get("queue_scope") == ["portland"]]
    assert {parse_qs(urlsplit(value).query)["workflow_tab"][0] for value in workflow_links} >= set(app.WORKFLOW_TABS)


def test_preview_count_is_separate_from_active_navigation_counts(client, monkeypatch, partition_pool):
    install_pool(monkeypatch, partition_pool)
    contexts = []

    def capture(sender, template, context, **extra):
        contexts.append(context)

    with template_rendered.connected_to(capture, app.app):
        response = client.get("/queue?queue_scope=portland")
    assert response.status_code == 200
    context = contexts[-1]
    assert context["queue_scope_counts"] == {"main": 3, "triage": 2, "portland": 0}
    assert context["portland_preview_count"] == 2
    assert sum(context["queue_scope_counts"].values()) == len(partition_pool)


@pytest.mark.parametrize("endpoint", ["review", "acknowledge"])
@pytest.mark.parametrize("valid_csrf", [False, True])
def test_review_and_ack_redirect_keep_portland_scope(client, monkeypatch, clean_state, endpoint, valid_csrf):
    install_pool(monkeypatch, [ticket()])
    seed(clean_state, 960001, "Resolved", reviewed_at=OLDER)
    client.get("/queue?queue_scope=portland")
    with client.session_transaction() as session:
        token = session["csrf_token"]
    response = client.post("/queue/api/" + endpoint, data={
        "ticket_id": "960001", "review_result": "Needs Follow-Up",
        "csrf_token": token if valid_csrf else "invalid",
        "queue_scope": "portland", "workflow_tab": "resolved", "days": "30",
    })
    assert response.status_code == 303
    query = parse_qs(urlsplit(response.headers["Location"]).query)
    assert query["queue_scope"] == ["portland"]
    assert query["workflow_tab"] == ["resolved"]
    assert query["days"] == ["30"]
    state = app.load_review_rows()[960001]
    if valid_csrf:
        assert state["review_result"] == ("Needs Follow-Up" if endpoint == "review" else "Resolved")
        assert state["reviewed_updated_at"] == STAMP
    else:
        assert state["review_result"] == "Resolved"
        assert state["reviewed_updated_at"] == OLDER


@pytest.mark.parametrize("enabled", [False, True])
def test_closed_mode_is_unsplit_and_keeps_normal_return(client, monkeypatch, enabled):
    activate(monkeypatch, enabled)
    pool = [ticket(960001, status=5, tags=[]), ticket(960002, order=OTHER, status=5, tags=[]), ticket(960003)]
    install_pool(monkeypatch, pool)
    args = MultiDict([("mode", "closed"), ("queue_scope", "portland"),
                      ("normal_queue_scope", "portland"), ("normal_workflow_tab", "followup")])
    parsed = app.filters_from_args(args)
    assert parsed["normal_return"]["queue_scope"] == "portland"
    assert parse_qs(app.filter_query_string(parsed))["normal_queue_scope"] == ["portland"]
    for scope in ("main", "triage", "portland"):
        response = client.get("/queue", query_string=list(args.items()) + [("queue_scope", scope)])
        assert response.status_code == 200
        assert html_ids(response) == [960001, 960002]
        assert b'aria-label="Queue scope"' not in response.data
    text = response.get_data(as_text=True)
    normal_keys = re.search(r"(?:var|const|let) normalKeys\s*=\s*\[([^]]+)\]", text)
    assert normal_keys and re.search(r"['\"]queue_scope['\"]", normal_keys.group(1))
    assert "normal_queue_scope" in unescape(text)


@pytest.mark.parametrize("enabled", [False, True])
def test_xlsx_order_matches_view_and_has_classification_notice(client, monkeypatch, clean_state, enabled):
    activate(monkeypatch, enabled)
    seed(clean_state, 960002, "Opened / In Review")
    before = snapshot(clean_state)
    pool = [ticket(960003, created_at=STAMP), ticket(960001),
            ticket(960002, order=PORTLAND_OTHER, created_at="2026-07-01T12:00:00Z"),
            ticket(960004, order=OTHER)]
    install_pool(monkeypatch, pool)
    expected = ids(app.build_current_queue_view(pool, config("portland")))
    html = client.get("/queue?queue_scope=portland")
    response = client.get("/queue/export.xlsx?queue_scope=portland")
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.content_type
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames[0] == ("Portland" if enabled else "Portland Preview")
    sheet = workbook.worksheets[0]
    exported_ids = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert exported_ids == expected == html_ids(html)
    assert "Classification Notice" in workbook.sheetnames
    notice = " ".join(str(cell.value) for row in workbook["Classification Notice"] for cell in row if cell.value is not None)
    assert "grammar" in notice.lower()
    if not enabled:
        assert "provisional" in notice.lower()
        assert "preview" in notice.lower()
    all_text = " ".join(str(cell.value) for ws in workbook for row in ws for cell in row if cell.value is not None)
    assert all(order not in all_text for order in (PORTLAND, PORTLAND_OTHER, OTHER))
    assert sheet.freeze_panes == "A2"
    assert snapshot(clean_state) == before


def test_empty_preview_export_has_notice_and_headers(client, monkeypatch):
    install_pool(monkeypatch, [ticket(order=OTHER)])
    response = client.get("/queue/export.xlsx?queue_scope=portland")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames[0] == "Portland Preview"
    assert workbook.worksheets[0].max_row == 1
    assert workbook.worksheets[0].cell(1, 1).value == "Ticket #"
    assert "Classification Notice" in workbook.sheetnames


@pytest.mark.parametrize("enabled", [False, True])
def test_read_paths_never_retrieve_read_keys_schedule_or_write_cache(client, monkeypatch, clean_state, enabled):
    activate(monkeypatch, enabled)
    install_pool(monkeypatch, [ticket()])
    seed(clean_state, 960001, "Resolved")
    before = snapshot(clean_state)

    def forbidden(*args, **kwargs):
        pytest.fail("read path attempted retrieval, credentials, scheduling, or cache mutation")

    for name in ("load_api_key", "fetch_live_queue", "paginate_tickets", "fetch_ticket_conversations",
                 "save_live_queue_cache", "_start_normal_queue_refresh", "start_auto_refresh_scheduler",
                 "initialize_live_auto_refresh", "reset_auto_refresh_countdown"):
        monkeypatch.setattr(app, name, forbidden)
    monkeypatch.setattr(app.queue_live.JOB, "start", forbidden)
    monkeypatch.setattr(app.closed_live.JOB, "start", forbidden)
    pool = [ticket()]
    original = copy.deepcopy(pool)
    for scope in ("main", "triage", "portland"):
        app.build_current_queue_view(pool, config(scope))
        for path in ("/queue", "/queue/export.xlsx"):
            assert client.get(path, query_string={"queue_scope": scope}).status_code == 200
    assert pool == original
    assert snapshot(clean_state) == before


def test_legacy_offline_population_unchanged_when_disabled(monkeypatch, partition_pool):
    # Emulate old metadata-poor fixtures without opening the fixture file.
    monkeypatch.delenv("FRESHDESK_OFFLINE_CACHE", raising=False)
    pool = copy.deepcopy(partition_pool)
    for value in pool:
        value.pop("group_id", None)
        value.pop("type", None)
    for scope in ("main", "triage"):
        view = app.build_current_queue_view(pool, config(scope))
        assert ids(view) == [value["id"] for value in pool]
        assert view[2] == {"main": len(pool), "triage": 0, "portland": 0}
    assert ids(app.build_current_queue_view(pool, config("portland"))) == [960001, 960002]


def test_activated_offline_population_uses_real_three_way_split(monkeypatch, partition_pool):
    monkeypatch.delenv("FRESHDESK_OFFLINE_CACHE", raising=False)
    activate(monkeypatch, True)
    expected = {"main": [960003, 960005], "triage": [960004], "portland": [960001, 960002]}
    for scope, wanted in expected.items():
        view = app.build_current_queue_view(partition_pool, config(scope))
        assert ids(view) == wanted
        assert view[2] == {"main": 2, "triage": 1, "portland": 2}


def test_activation_and_rollback_only_reclassify_existing_review_rows(clean_state, monkeypatch):
    seed(clean_state, 960001, "Resolved")
    before = snapshot(clean_state)
    value = ticket()
    for enabled, expected_scope in ((False, "main"), (True, "portland"), (False, "main")):
        activate(monkeypatch, enabled)
        for scope in ("main", "triage", "portland"):
            view = app.build_current_queue_view([value], config(scope, workflow_tab="resolved"))
            # The disabled Portland preview overlaps Main deliberately.
            expected = scope == expected_scope or (not enabled and scope == "portland")
            assert ids(view) == ([960001] if expected else [])
            assert sum(view[2].values()) == 1
        assert snapshot(clean_state) == before


def test_changed_or_removed_order_evidence_reclassifies_without_review_writes(clean_state, monkeypatch):
    activate(monkeypatch, True)
    seed(clean_state, 960001, "No Action", reviewed_at=OLDER)
    before = snapshot(clean_state)
    for order, expected in ((PORTLAND, "portland"), (OTHER, "main"), (None, "main"),
                            (PORTLAND + "," + OTHER, "main")):
        value = ticket(order=order)
        assert app.queue_scope_for_ticket(value, portland_enabled=True) == expected
        assert ids(app.build_current_queue_view([value], config(expected, workflow_tab="no_action"))) == [960001]
        assert snapshot(clean_state) == before


def test_activated_metadata_poor_offline_tickets_do_not_leak_into_other_scopes(monkeypatch):
    monkeypatch.delenv("FRESHDESK_OFFLINE_CACHE", raising=False)
    activate(monkeypatch, True)
    pool = [ticket(960001), ticket(960002, order=OTHER)]
    for value in pool:
        value.pop("group_id")
        value.pop("type")
    assert ids(app.build_current_queue_view(pool, config("main"))) == []
    assert ids(app.build_current_queue_view(pool, config("triage"))) == [960002]
    assert ids(app.build_current_queue_view(pool, config("portland"))) == [960001]
